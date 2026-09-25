from __future__ import annotations

import base64
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from cureflow_app.server import (
    ApiInputError,
    export_repaired_copy,
    inspect_tabular_input,
    preview_repaired_copy,
    xlsx_base64_to_tables,
)
from workflow.repair_analysis import (
    apply_confirmed_fixes_to_csv,
    apply_safe_fixes_to_csv,
    approved_manual_corrections,
    approved_safe_fixes,
    build_repair_plan,
)


def _plan(csv_text: str, **kwargs):
    return build_repair_plan(csv_text, filename="input.csv", **kwargs)


def test_repair_plan_locates_safe_whitespace_fix() -> None:
    plan = _plan("姓名,成绩\n 张三 ,88\n李四,92\n")
    issue = next(item for item in plan["issues"] if item["type"] == "trim_whitespace")
    assert issue["location"]["cell"] == "A2"
    assert issue["originalValue"] == " 张三 "
    assert issue["proposedValue"] == "张三"
    assert issue["autoApplicable"] is True
    assert plan["summary"]["safeFixCount"] == 1


def test_repair_plan_counts_all_issues_and_pages_beyond_first_250() -> None:
    source = "编号,备注\n" + "".join(f"R{index}, 值{index} \n" for index in range(270))
    first = _plan(source)
    later = _plan(source, issue_offset=250, issue_limit=40)
    assert first["summary"]["issueCount"] == 270
    assert first["summary"]["shownIssueCount"] == 250
    assert first["summary"]["byType"] == {"trim_whitespace": 270}
    assert first["summary"]["byColumn"] == {"备注": 270}
    assert later["summary"]["matchingIssueCount"] == 270
    assert len(later["issues"]) == 20
    assert later["issues"][0]["location"]["cell"] == "B252"


def test_repair_plan_filters_on_full_population_not_only_first_page() -> None:
    source = "编号,备注\n" + "".join(f"R{index}, 值{index} \n" for index in range(260))
    plan = _plan(source, issue_offset=250, issue_limit=10, issue_type_filter="trim_whitespace", column_filter="备注")
    assert [item["location"]["cell"] for item in plan["issues"]] == [f"B{row}" for row in range(252, 262)]
    assert plan["summary"]["matchingIssueCount"] == 260


def test_repair_plan_preserves_real_xlsx_row_numbers() -> None:
    plan = _plan("编号,备注\nA, ok \nB,clean\n", sheet="主表", source_row_numbers=[5, 9])
    issue = next(item for item in plan["issues"] if item["type"] == "trim_whitespace")
    assert issue["location"] == {
        "sheet": "主表", "row": 5, "column": "备注", "cell": "B5", "dataIndex": 0,
    }


def test_missing_value_is_review_only_and_never_guessed() -> None:
    plan = _plan("编号,成绩\nA,90\nB,\nC,80\n")
    issue = next(item for item in plan["issues"] if item["type"] == "missing_value")
    assert issue["location"]["cell"] == "B3"
    assert issue["proposedValue"] is None
    assert issue["autoApplicable"] is False


def test_duplicate_row_requires_review() -> None:
    plan = _plan("编号,值\nA,1\nA,1\n")
    issue = next(item for item in plan["issues"] if item["type"] == "duplicate_row")
    assert issue["autoApplicable"] is False
    assert issue["evidence"]["duplicatesSourceRow"] == 2


def test_score_and_rank_bounds_are_located() -> None:
    plan = _plan("学号,综合成绩,班级排名,班级总人数\nA,980,12,10\n")
    types = {item["type"]: item for item in plan["issues"]}
    assert types["score_out_of_range"]["location"]["cell"] == "B2"
    assert types["rank_out_of_range"]["location"]["cell"] == "C2"
    assert all(not item["autoApplicable"] for item in types.values())


def test_iqr_extreme_is_exploration_only_not_a_repair_issue() -> None:
    values = [*range(1, 31), 1000]
    csv_text = "样本,测量值\n" + "\n".join(f"S{i},{value}" for i, value in enumerate(values)) + "\n"
    plan = _plan(csv_text)
    assert plan["issues"] == []
    assert plan["summary"]["issueCount"] == 0
    screening = next(item for item in plan["summary"]["iqrScreening"] if item["column"] == "测量值")
    assert screening["status"] == "applied"
    assert screening["candidateCount"] == 1
    assert screening["upperBound"] < 1000
    assert "不是修正问题" in screening["reason"]


def test_derived_rate_is_not_iqr_screened_or_queued_for_correction() -> None:
    values = [round((index % 40 - 20) / 100, 3) for index in range(100)] + [-0.71]
    csv_text = "学业进步率,实测值\n" + "\n".join(f"{rate},{index + 1}" for index, rate in enumerate(values)) + "\n"
    plan = _plan(csv_text)
    rate_screen = next(item for item in plan["summary"]["iqrScreening"] if item["column"] == "学业进步率")
    assert rate_screen["status"] == "skipped"
    assert "分子、分母" in rate_screen["reason"]
    assert not any(item["location"]["column"] == "学业进步率" for item in plan["issues"])


def test_iqr_skips_small_or_discrete_numeric_columns() -> None:
    small = _plan("测量值\n" + "\n".join(str(value) for value in [*range(1, 8), 1000]) + "\n")
    assert small["summary"]["iqrScreening"][0]["status"] == "skipped"
    assert "有效数值" in small["summary"]["iqrScreening"][0]["reason"]
    discrete = _plan("等级\n" + "\n".join(str(value) for value in [1, 2, 3] * 10 + [99]) + "\n")
    assert not any(item["type"] == "statistical_outlier" for item in discrete["issues"])
    assert "不同数值过少" in discrete["summary"]["iqrScreening"][0]["reason"]


def test_iqr_does_not_mix_explicit_groups_or_time_series() -> None:
    grouped = "组别,测量值\n" + "\n".join(
        f"{group},{value}" for group, values in (("A", range(1, 21)), ("B", range(101, 121))) for value in values
    ) + "\n"
    plan = _plan(grouped)
    measure = next(item for item in plan["summary"]["iqrScreening"] if item["column"] == "测量值")
    assert measure["status"] == "skipped"
    assert "分组字段" in measure["reason"]
    assert not any(item["type"] == "statistical_outlier" for item in plan["issues"])

    timed = "时间,测量值\n" + "\n".join(f"T{i},{i if i < 30 else 1000}" for i in range(31)) + "\n"
    time_plan = _plan(timed)
    timed_measure = next(item for item in time_plan["summary"]["iqrScreening"] if item["column"] == "测量值")
    assert timed_measure["status"] == "skipped"
    assert "时间字段" in timed_measure["reason"]

    selected_group = "性别,测量值\n" + "\n".join(
        f"{group},{value}" for group, values in (("甲", range(1, 21)), ("乙", range(101, 121))) for value in values
    ) + "\n"
    selected_plan = _plan(selected_group, analysis_config={"groupColumn": "性别"})
    selected_measure = next(item for item in selected_plan["summary"]["iqrScreening"] if item["column"] == "测量值")
    assert selected_measure["status"] == "skipped"
    assert "本次已指定分组字段“性别”" in selected_measure["reason"]


def test_iqr_skips_mixed_distribution_with_many_extremes() -> None:
    values = [*range(1, 96), *range(1000, 1010)]
    plan = _plan("测量值\n" + "\n".join(map(str, values)) + "\n")
    assessment = plan["summary"]["iqrScreening"][0]
    assert assessment["status"] == "skipped"
    assert "集中出现" in assessment["reason"]
    assert not any(item["type"] == "statistical_outlier" for item in plan["issues"])


def test_runtime_duration_is_not_mistaken_for_time_axis() -> None:
    values = [*range(1, 31), 1000]
    plan = _plan("runtime\n" + "\n".join(map(str, values)) + "\n")
    assert plan["summary"]["iqrScreening"][0]["status"] == "applied"
    assert plan["summary"]["iqrScreening"][0]["candidateCount"] == 1
    assert not plan["issues"]


def test_only_safe_approved_fixes_can_be_applied() -> None:
    plan = _plan("编号,备注\nA, ok \nB,\n")
    safe = next(item for item in plan["issues"] if item["autoApplicable"])
    manual = next(item for item in plan["issues"] if not item["autoApplicable"])
    assert approved_safe_fixes(plan, [safe["id"]]) == [safe]
    with pytest.raises(ValueError, match="人工复核"):
        approved_safe_fixes(plan, [manual["id"]])
    with pytest.raises(ValueError, match="不存在"):
        approved_safe_fixes(plan, ["fix_missing"])


def test_apply_safe_fix_creates_new_csv_content() -> None:
    source = "编号,备注\nA, ok \n"
    plan = _plan(source)
    safe = next(item for item in plan["issues"] if item["autoApplicable"])
    repaired = apply_safe_fixes_to_csv(source, [safe])
    assert repaired == "编号,备注\nA,ok\n"
    assert source == "编号,备注\nA, ok \n"


def test_export_repaired_xlsx_keeps_original_bytes_and_preserves_other_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "主表"
    sheet.append(["编号", "备注", "公式"])
    sheet.append(["A", " ok ", "=1+1"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    original = stream.getvalue()
    encoded = base64.b64encode(original).decode("ascii")
    inspection = inspect_tabular_input({"filename": "待修.xlsx", "xlsxBase64": encoded})
    plan = build_repair_plan(
        xlsx_base64_to_tables(encoded, "待修.xlsx")[0]["csvText"],
        filename="待修.xlsx", sheet="主表", source_row_numbers=[2],
    )
    safe = next(item for item in plan["issues"] if item["autoApplicable"])
    result = export_repaired_copy({
        "filename": "待修.xlsx", "xlsxBase64": encoded,
        "selectedSheet": inspection["selectedSheet"], "approvedFixIds": [safe["id"]],
    })
    assert base64.b64decode(encoded) == original
    repaired = load_workbook(BytesIO(base64.b64decode(result["xlsxBase64"])), data_only=False)
    assert repaired["主表"]["B2"].value == "ok"
    assert repaired["主表"]["C2"].value == "=1+1"
    repaired.close()
    assert result["audit"]["sourceFileModified"] is False
    assert result["audit"]["manualReviewItemsApplied"] == 0


def test_export_rejects_manual_review_issue() -> None:
    source = "编号,成绩\nA,999\n"
    plan = _plan(source)
    manual = next(item for item in plan["issues"] if not item["autoApplicable"])
    with pytest.raises(ApiInputError, match="人工复核"):
        export_repaired_copy({
            "filename": "input.csv", "csvText": source,
            "approvedFixIds": [manual["id"]],
        })


def test_manual_correction_requires_explicit_confirmation() -> None:
    plan = _plan("编号,成绩\nA,999\n")
    issue = next(item for item in plan["issues"] if item["type"] == "score_out_of_range")
    with pytest.raises(ValueError, match="明确确认"):
        approved_manual_corrections(
            plan, [{"issueId": issue["id"], "newValue": "99"}], confirmed=False,
        )


def test_manual_correction_applies_exact_user_value_to_new_csv() -> None:
    source = "编号,成绩\nA,999\n"
    plan = _plan(source)
    issue = next(item for item in plan["issues"] if item["type"] == "score_out_of_range")
    fixes = approved_manual_corrections(
        plan, [{"issueId": issue["id"], "newValue": "99"}], confirmed=True,
    )
    assert apply_confirmed_fixes_to_csv(source, fixes) == "编号,成绩\nA,99\n"


def test_manual_correction_rejects_duplicate_row_and_formula_input() -> None:
    duplicate_plan = _plan("编号,成绩\nA,90\nA,90\n")
    duplicate = next(item for item in duplicate_plan["issues"] if item["type"] == "duplicate_row")
    with pytest.raises(ValueError, match="不能通过单元格替换"):
        approved_manual_corrections(
            duplicate_plan, [{"issueId": duplicate["id"], "newValue": "B"}], confirmed=True,
        )
    score_plan = _plan("编号,成绩\nA,999\n")
    score = next(item for item in score_plan["issues"] if item["type"] == "score_out_of_range")
    with pytest.raises(ValueError, match="公式注入"):
        approved_manual_corrections(
            score_plan, [{"issueId": score["id"], "newValue": "=1+1"}], confirmed=True,
        )


def test_export_accepts_confirmed_manual_correction_and_audits_it() -> None:
    source = "编号,成绩\nA,999\n"
    plan = _plan(source)
    issue = next(item for item in plan["issues"] if item["type"] == "score_out_of_range")
    result = export_repaired_copy({
        "filename": "input.csv",
        "csvText": source,
        "approvedFixIds": [],
        "manualCorrections": [{"issueId": issue["id"], "newValue": "99"}],
        "confirmManualCorrections": True,
    })
    assert result["csvText"] == "编号,成绩\nA,99\n"
    assert result["audit"]["safeFixesApplied"] == 0
    assert result["audit"]["manualReviewItemsApplied"] == 1
    assert result["audit"]["manualCorrectionIds"] == [issue["id"]]


def test_repair_preview_compares_before_after_without_modifying_source() -> None:
    source = "编号,成绩\nA,999\nB,80\nC,90\nD,85\n"
    plan = _plan(source)
    issue = next(item for item in plan["issues"] if item["type"] == "score_out_of_range")
    preview = preview_repaired_copy({
        "filename": "scores.csv",
        "csvText": source,
        "approvedFixIds": [],
        "manualCorrections": [{"issueId": issue["id"], "newValue": "99"}],
        "confirmManualCorrections": True,
    })
    assert preview["providerUsed"] is False
    assert preview["sourceFileModified"] is False
    assert preview["summary"] == {
        "beforeIssueCount": 1,
        "afterIssueCount": 0,
        "resolvedIssueCount": 1,
        "remainingIssueCount": 0,
        "introducedIssueCount": 0,
        "changedCellCount": 1,
    }
    assert preview["changedCells"][0]["originalValue"] == "999"
    assert preview["changedCells"][0]["correctedValue"] == "99"
    assert preview["changedCells"][0]["source"] == "user_confirmed"
    assert source == "编号,成绩\nA,999\nB,80\nC,90\nD,85\n"
