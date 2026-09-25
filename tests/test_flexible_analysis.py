"""Adaptive schema detection and generic deterministic table profiling."""
from __future__ import annotations

import base64
import json
import re
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook

from workflow.flexible_analysis import (
    analyze_generic_table,
    analyze_generic_workbook,
    prepare_tabular_csv,
)
from workflow.repair_analysis import build_repair_plan
from cureflow_app.server import make_handler, probe_provider_connection
from cureflow_app.provider_runtime import (
    DesktopProviderError,
    DesktopProviderRuntime,
    ProviderRun,
)
from llm.base import LLMResponse, UsageType


class GroundedInterpretationClient:
    def __init__(self, *, invalid_ref: bool = False) -> None:
        self.invalid_ref = invalid_ref
        self.prompts: list[str] = []

    def generate(self, prompt: str, *, task_name: str = "") -> LLMResponse:
        self.prompts.append(prompt)
        refs = re.findall(
            r"desktop_[A-Za-z0-9_]+/metric_[A-Za-z0-9_]+@v\d+", prompt
        )
        reference = "fabricated/ref@v1" if self.invalid_ref else refs[0]
        payload = {
            "summary": "数据完整性和异常检查结果需要结合字段用途复核。",
            "findings": [{
                "title": "先核对数据质量",
                "explanation": "现有汇总指标提示应优先检查缺失与异常记录。",
                "confidence": "high",
                "evidenceRefs": [reference],
            }],
            "actions": [{
                "action": "核对缺失记录",
                "reason": "补齐数据后再进行进一步分析。",
                "evidenceRefs": [reference],
            }],
            "limitations": ["未接收原始数据行，不能判断具体记录。"],
        }
        return LLMResponse(
            text=json.dumps(payload, ensure_ascii=False),
            input_tokens=120,
            output_tokens=60,
            latency_ms=25.0,
            provider="openai",
            model="gpt-5.6-luna",
            usage_type=UsageType.PROVIDER_REPORTED,
            request_id="resp_private_test",
        )


class CompactClaimInterpretationClient(GroundedInterpretationClient):
    def generate(self, prompt: str, *, task_name: str = "") -> LLMResponse:
        self.prompts.append(prompt)
        reference = re.findall(
            r"desktop_[A-Za-z0-9_]+/metric_[A-Za-z0-9_]+@v\d+", prompt
        )[0]
        return LLMResponse(
            text=json.dumps({
                "claim_text": "当前数据量有限，应先核对缺失情况后再扩大分析。",
                "evidence_refs": [reference],
                "confidence": 0.82,
            }, ensure_ascii=False),
            input_tokens=80,
            output_tokens=30,
            latency_ms=20.0,
            provider="deepseek",
            model="deepseek-v4-flash",
            usage_type=UsageType.PROVIDER_REPORTED,
        )


class BlankSummaryCompactClaimClient(CompactClaimInterpretationClient):
    def __init__(self, blank_summary: object) -> None:
        super().__init__()
        self.blank_summary = blank_summary

    def generate(self, prompt: str, *, task_name: str = "") -> LLMResponse:
        response = super().generate(prompt, task_name=task_name)
        payload = json.loads(response.text)
        payload["summary"] = self.blank_summary
        return LLMResponse(
            text=json.dumps(payload, ensure_ascii=False),
            input_tokens=response.input_tokens,
            output_tokens=response.output_tokens,
            latency_ms=response.latency_ms,
            provider=response.provider,
            model=response.model,
            usage_type=response.usage_type,
        )


class ProviderErrorInterpretationClient(GroundedInterpretationClient):
    def generate(self, prompt: str, *, task_name: str = "") -> LLMResponse:
        self.prompts.append(prompt)
        return LLMResponse(
            text=json.dumps({
                "error": "AUTH_ERROR",
                "message": "HTTP 401: invalid API key",
            }),
            provider="openai",
            model="gpt-5.6-luna",
            usage_type=UsageType.UNAVAILABLE,
        )


class ConnectionProbeClient:
    def __init__(self, *, auth_error: bool = False) -> None:
        self.auth_error = auth_error

    def generate(self, prompt: str, *, task_name: str = "") -> LLMResponse:
        assert "不包含用户文件或业务数据" in prompt
        assert task_name == "provider_connection_probe"
        payload = (
            {
                "error": "AUTH_ERROR",
                "message": "HTTP 401 incorrect_api_key sk-ws-********secret",
            }
            if self.auth_error
            else {
                "claim_text": "连接验证成功",
                "evidence_refs": ["connection_probe/fixed@v1"],
                "confidence": 1.0,
            }
        )
        return LLMResponse(
            text=json.dumps(payload, ensure_ascii=False),
            provider="openai",
            model="gpt-5.6-luna",
            usage_type=(
                UsageType.UNAVAILABLE
                if self.auth_error else UsageType.PROVIDER_REPORTED
            ),
        )


def _provider_run(client: GroundedInterpretationClient) -> ProviderRun:
    return ProviderRun(
        mode="OPENAI_RESPONSES_LIVE",
        provider="openai",
        model="gpt-5.6-luna",
        live=True,
        client=client,
        call_budget=6,
    )


def _generic_csv() -> str:
    return (
        "Time,Surface Displacement (mm),Rainfall (mm),Station\n"
        "2025-05-01 16:40,0.372,0,A\n"
        "2025-05-01 16:50,0.369,,A\n"
        "2025-05-01 17:00,8.500,2,B\n"
    )


def _workbook_tables() -> list[dict[str, str]]:
    return [
        {
            "sheet": "训练集",
            "csvText": "时间,目标值,降雨量,孔隙水压力\n2024-01-01,1.2,0,30\n2024-01-02,1.4,2,31\n",
        },
        {
            "sheet": "实验集",
            "csvText": "时间,阶段标签,目标值,降雨量,孔隙水压力\n2025-01-01,1,,0,29\n2025-01-02,2,,3,32\n",
        },
    ]


def _workbook_base64() -> str:
    workbook = Workbook()
    training = workbook.active
    training.title = "训练集"
    for row in [["时间", "目标值", "降雨量", "孔隙水压力"], ["2024-01-01", 1.2, 0, 30], ["2024-01-02", 1.4, 2, 31]]:
        training.append(row)
    experiment = workbook.create_sheet("实验集")
    for row in [["时间", "阶段标签", "目标值", "降雨量", "孔隙水压力"], ["2025-01-01", 1, None, 0, 29], ["2025-01-02", 2, None, 3, 32]]:
        experiment.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return base64.b64encode(stream.getvalue()).decode("ascii")


@pytest.fixture
def handler_request(tmp_path: Path):
    handler_type = make_handler(tmp_path)

    def invoke(payload: object):
        body = json.dumps(payload).encode("utf-8")
        handler = object.__new__(handler_type)
        handler.path = "/api/analyze"
        handler.headers = {"Content-Length": str(len(body))}
        handler.rfile = BytesIO(body)
        handler.wfile = BytesIO()
        response: dict[str, object] = {"status": None, "headers": {}}
        handler.send_response = lambda status: response.update(status=status)
        handler.send_header = lambda key, value: response["headers"].__setitem__(key, value)
        handler.end_headers = lambda: None
        handler.do_POST()
        response["json"] = json.loads(handler.wfile.getvalue().decode("utf-8"))
        return response

    return invoke


def test_business_sample_keeps_established_mode() -> None:
    sample = Path(__file__).resolve().parents[1] / "data" / "sample_company.csv"
    normalized, metadata = prepare_tabular_csv(sample.read_text(encoding="utf-8"))
    assert metadata["analysisMode"] == "business"
    assert metadata["missingBusinessFields"] == []
    assert set(pd.read_csv(StringIO(normalized)).columns) >= {
        "year", "month", "revenue", "cost", "users", "orders", "visits"
    }


def _business_goal_input(file_format: str = "csv") -> dict:
    headers = ["年份", "月份", "营业收入", "运营成本", "月活用户", "订单数", "访问量", "渠道", "日期", "客户编号"]
    rows = [
        [2025, i, 100 + i * 10, 50 + i * 3, 10 + i, 5 + i, 30 + i * 2,
         "甲" if i <= 4 else "乙", f"2025-{i:02d}-01", f"00{i}"]
        for i in range(1, 9)
    ]
    if file_format == "xlsx":
        book = Workbook()
        book.active.title = "业务明细"
        for row in [headers, *rows]:
            book.active.append(row)
        stream = BytesIO()
        book.save(stream)
        book.close()
        return {"filename": "业务明细.xlsx", "xlsxBase64": base64.b64encode(stream.getvalue()).decode("ascii"), "selectedSheet": "业务明细"}
    return {
        "filename": "业务明细.csv",
        "csvText": "\n".join(",".join(map(str, row)) for row in [headers, *rows]) + "\n",
    }


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
@pytest.mark.parametrize("goal", [
    "data_quality", "descriptive_profile", "group_comparison",
    "relationship_review", "trend_forecast", "causal_evaluation",
])
def test_explicit_goal_keeps_original_business_columns_and_reaches_interpreter(
    handler_request, file_format: str, goal: str,
) -> None:
    from workflow.analysis_tools import GOAL_TOOLS

    client = GroundedInterpretationClient()
    config = {"measureColumn": "营业收入", "groupColumn": "渠道", "timeColumn": "日期", "identifierColumn": "客户编号"}
    with patch("cureflow_app.server.DesktopProviderRuntime.resolve", return_value=_provider_run(client)), patch(
        "cureflow_app.server._analyze_csv_with_provider",
        side_effect=AssertionError("explicit goal must not execute legacy business claims"),
    ) as legacy:
        response = handler_request({
            **_business_goal_input(file_format), "analysisGoal": goal, "analysisConfig": config,
            "providerMode": "openai_responses", "confirmLive": True, "confirmMaxCalls": 6,
        })
    assert response["status"] == 200, response["json"]
    legacy.assert_not_called()
    result = response["json"]
    assert result["input"]["analysisConfig"] == config
    assert result["input"]["rowCount"] == 8
    assert result["input"]["columnCount"] == 10
    assert set(tool["key"] for tool in result["analysisTools"]["executedTools"]) == GOAL_TOOLS[goal]
    assert len(client.prompts) == 1
    packet = json.loads(client.prompts[0].rsplit("<UNTRUSTED_EVIDENCE>", 1)[1].split("</UNTRUSTED_EVIDENCE>", 1)[0])
    assert packet["analysisTools"]["goal"] == goal
    assert packet["analysisTools"]["analysisConfig"] == config
    questions = {
        "data_quality": "数据可靠吗？", "descriptive_profile": "现在是什么情况？",
        "group_comparison": "哪些群体有差异？", "relationship_review": "哪些因素有关联？",
        "trend_forecast": "趋势怎样，会怎么变？", "causal_evaluation": "措施真的有效吗？",
    }
    assert result["report"].startswith(f"# {questions[goal]}")
    if file_format == "xlsx":
        assert result["input"]["sheet"] == "业务明细"
    if goal == "relationship_review":
        candidates = result["analysisTools"]["executedTools"][0]["payload"]["candidates"]
        assert candidates
        assert all("营业收入" in {item["left"], item["right"]} for item in candidates)
    if goal == "trend_forecast":
        tool = next(item for item in result["analysisTools"]["executedTools"] if item["key"] == "relationship_screen")
        assert tool["payload"]["timeTrend"]["timeColumn"] == "日期"
        assert tool["payload"]["timeTrend"]["comparableRows"] == 8
    assert result["cost"]["llmCalls"] == 1
    assert "revenue_growth" not in {metric["key"] for metric in result["metrics"]}


def test_invalid_business_goal_measure_is_rejected_before_provider_resolution(handler_request) -> None:
    with patch("cureflow_app.server.DesktopProviderRuntime.resolve") as resolve:
        response = handler_request({
            **_business_goal_input(), "analysisGoal": "relationship_review",
            "analysisConfig": {"measureColumn": "渠道"},
            "providerMode": "openai_responses", "confirmLive": True, "confirmMaxCalls": 6,
        })
    assert response["status"] == 400
    assert "主指标" in response["json"]["error"]
    resolve.assert_not_called()


def test_chinese_business_headers_are_mapped_without_fuzzy_guessing() -> None:
    source = (
        "年份,月份,营业收入,运营成本,月活用户,订单数,访问量\n"
        "2024,1,100,60,10,5,20\n2025,1,150,70,12,8,25\n"
    )
    normalized, metadata = prepare_tabular_csv(source)
    assert metadata["analysisMode"] == "business"
    assert metadata["columnMapping"]["revenue"] == "营业收入"
    assert list(pd.read_csv(StringIO(normalized)).columns) == [
        "year", "month", "revenue", "cost", "users", "orders", "visits"
    ]


def test_date_column_can_derive_year_and_month() -> None:
    source = (
        "日期,营收,成本,用户数,订单数,访问量\n"
        "2024-01-01,100,60,10,5,20\n2025-01-01,150,70,12,8,25\n"
    )
    normalized, metadata = prepare_tabular_csv(source)
    frame = pd.read_csv(StringIO(normalized))
    assert metadata["analysisMode"] == "business"
    assert metadata["derivedColumns"] == ["year", "month"]
    assert list(frame["year"]) == [2024, 2025]


def test_unrelated_monitoring_columns_use_generic_mode() -> None:
    _, metadata = prepare_tabular_csv(_generic_csv())
    assert metadata["analysisMode"] == "generic_profile"
    assert "revenue" in metadata["missingBusinessFields"]
    assert metadata["columnMapping"] == {}


def test_generic_profile_returns_six_truthful_metrics(tmp_path: Path) -> None:
    result = analyze_generic_table(_generic_csv(), "monitoring.csv", tmp_path)
    metrics = {item["key"]: item for item in result["metrics"]}
    assert set(metrics) == {
        "row_count", "column_count", "completeness", "numeric_columns",
        "duplicate_rows", "outlier_rate",
    }
    assert metrics["row_count"]["value"] == 3
    assert metrics["column_count"]["value"] == 4
    assert metrics["completeness"]["display"] == "91.67%"
    assert result["mode"] == "OFFLINE_DETERMINISTIC_PROFILE"
    assert result["cost"]["llmCalls"] == 0


def test_product_profile_does_not_report_skipped_iqr_as_zero_errors(tmp_path: Path) -> None:
    csv_text = "组别,测量值\n" + "\n".join(
        f"{group},{value}" for group, values in (("A", range(1, 21)), ("B", range(101, 121))) for value in values
    ) + "\n"
    plan = build_repair_plan(csv_text, filename="grouped.csv")
    result = analyze_generic_table(csv_text, "grouped.csv", tmp_path, repair_plan=plan)
    metric = next(item for item in result["metrics"] if item["key"] == "outlier_rate")
    assert metric["display"] == "未启用"
    assert "没有符合保守条件的 IQR 筛查列" in result["claims"][2]["content"]
    assert "不能把未筛查写成异常值为零" in result["claims"][2]["content"]
    assert "IQR 极端值复核：未启用" in result["report"]


def test_generic_profile_claims_have_persisted_evidence(tmp_path: Path) -> None:
    result = analyze_generic_table(_generic_csv(), "monitoring.csv", tmp_path)
    assert all(claim["grounding"]["verdict"] == "grounded" for claim in result["claims"])
    assert all(claim["evidenceRefs"] for claim in result["claims"])
    run_dir = tmp_path / "artifacts" / result["runId"]
    assert (run_dir / "raw_data__v1.json").exists()
    assert (run_dir / "field_column_profiles__v1.json").exists()
    assert len(list(run_dir.glob("claim_*__v1.json"))) == 3


def test_generic_report_names_source_columns_without_business_relabeling(tmp_path: Path) -> None:
    result = analyze_generic_table(_generic_csv(), "monitoring.csv", tmp_path)
    assert "Surface Displacement (mm)" in result["report"]
    assert "Rainfall (mm)" in result["report"]
    assert "营收增长率" not in result["report"]
    assert "不把任意字段强行解释" in result["disclaimer"]


def test_one_row_profile_marks_inference_as_insufficient(tmp_path: Path) -> None:
    result = analyze_generic_table("学号,姓名,专业排名\n1001,张三,1\n", "student.csv", tmp_path)
    assert result["dataSufficiency"]["level"] == "insufficient"
    assert result["dataSufficiency"]["label"] == "证据不足"
    assert "不能据此判断分布、趋势、相关性或异常模式" in result["dataSufficiency"]["summary"]
    assert "IQR 异常值：不可评估" in result["report"]
    assert next(item for item in result["metrics"] if item["key"] == "outlier_rate")["display"] == "不可评估"


def test_deterministic_profile_does_not_claim_zero_hallucination_risk(tmp_path: Path) -> None:
    result = analyze_generic_table(_generic_csv(), "monitoring.csv", tmp_path)
    assert result["cost"]["llmCalls"] == 0
    assert result["groundingSummary"]["observed_risk_rate_pct"] is None
    assert "幻觉风险率不适用" in result["groundingSummary"]["interpretation"]


def test_generic_profile_classifies_identifiers_ranks_counts_and_scores(tmp_path: Path) -> None:
    source = (
        "学号,姓名,综合成绩,专业排名,专业总人数,德育等级\n"
        + "\n".join(
            f"92513083{index:04d},学生{index},{70 + index / 10},{index},40,合格"
            for index in range(1, 41)
        )
        + "\n"
    )
    result = analyze_generic_table(source, "scores.csv", tmp_path)
    assert "| 学号 | 标识符 |" in result["report"]
    assert "| 综合成绩 | 连续成绩 |" in result["report"]
    assert "| 专业排名 | 顺序排名 |" in result["report"]
    assert "| 专业总人数 | 计数 |" in result["report"]
    assert "排名、计数和标识符不参与异常值判断" in result["claims"][2]["content"]
    assert result["semanticSummary"]["identifiers"] == [
        {"column": "学号", "missingCount": 0, "duplicateCount": 0}
    ]
    assert result["semanticSummary"]["scoreFields"][0]["column"] == "综合成绩"
    assert result["semanticSummary"]["rankChecks"] == [
        {
            "rankColumn": "专业排名",
            "totalColumn": "专业总人数",
            "checkedRows": 40,
            "outOfRangeRows": 0,
        }
    ]
    assert "## 可操作检测" in result["report"]


def test_api_automatically_routes_monitoring_csv_to_generic_profile(handler_request) -> None:
    response = handler_request({"filename": "monitoring.csv", "csvText": _generic_csv()})
    assert response["status"] == 200
    payload = response["json"]
    assert payload["input"]["analysisMode"] == "generic_profile"
    assert payload["input"]["analysisModeLabel"] == "通用表格分析"
    assert payload["provider"]["name"] == "local"
    assert payload["stages"][1]["key"] == "schema"


def test_generic_mode_uses_selected_provider_for_grounded_interpretation(
    handler_request,
) -> None:
    client = GroundedInterpretationClient()
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=_provider_run(client),
    ) as resolve:
        response = handler_request(
            {
                "filename": "monitoring.csv",
                "csvText": _generic_csv(),
                "providerMode": "openai_responses",
                "confirmLive": True,
                "confirmMaxCalls": 6,
            }
        )
    resolve.assert_called_once()
    assert response["status"] == 200
    payload = response["json"]
    assert payload["input"]["cloudProviderUsed"] is True
    assert payload["cost"]["llmCalls"] == 1
    assert payload["provider"]["name"] == "openai"
    assert payload["aiInterpretation"]["reviewStatus"] == "requires_human_review"
    assert payload["aiInterpretation"]["rawRowsSent"] is False
    assert payload["modeLabel"] == "OpenAI 已真实调用 · 证据解读"
    assert payload["report"].startswith("# 这份数据有哪些值得关注的问题？")
    assert "## 一句话看懂" in payload["report"]
    assert "## 建议下一步" in payload["report"]
    assert "@v1" not in payload["report"]
    assert "analysis_tool_" not in payload["report"]
    assert "2025-05-01" not in client.prompts[0]
    assert "0.372" not in client.prompts[0]
    assert '"topValues"' not in client.prompts[0]


def test_generic_mode_rejects_fabricated_evidence_reference(handler_request) -> None:
    client = GroundedInterpretationClient(invalid_ref=True)
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=_provider_run(client),
    ):
        response = handler_request({
            "filename": "monitoring.csv",
            "csvText": _generic_csv(),
            "providerMode": "openai_responses",
            "confirmLive": True,
            "confirmMaxCalls": 6,
        })
    assert response["status"] == 400
    assert "不存在的证据" in response["json"]["error"]


def test_generic_mode_accepts_frozen_compact_claim_contract(handler_request) -> None:
    client = CompactClaimInterpretationClient()
    provider_run = ProviderRun(
        mode="DEEPSEEK_CHAT_LIVE",
        provider="deepseek",
        model="deepseek-v4-flash",
        live=True,
        client=client,
        call_budget=6,
    )
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=provider_run,
    ):
        response = handler_request({
            "filename": "monitoring.csv",
            "csvText": _generic_csv(),
            "providerMode": "deepseek_chat",
            "confirmLive": True,
            "confirmMaxCalls": 6,
        })
    assert response["status"] == 200
    interpretation = response["json"]["aiInterpretation"]
    assert interpretation["summary"].startswith("当前数据量有限")
    assert interpretation["findings"][0]["confidence"] == "high"
    assert interpretation["findings"][0]["evidenceRefs"]
    assert interpretation["actions"] == []


@pytest.mark.parametrize("blank_summary", [None, "", "   "])
def test_generic_mode_normalizes_compact_contract_with_blank_summary(
    handler_request, blank_summary: object,
) -> None:
    client = BlankSummaryCompactClaimClient(blank_summary)
    provider_run = ProviderRun(
        mode="OPENAI_RESPONSES_LIVE",
        provider="openai",
        model="gpt-5.6-luna",
        live=True,
        client=client,
        call_budget=6,
    )
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=provider_run,
    ):
        response = handler_request({
            "filename": "monitoring.csv",
            "csvText": _generic_csv(),
            "providerMode": "openai_responses",
            "confirmLive": True,
            "confirmMaxCalls": 6,
        })
    assert response["status"] == 200
    interpretation = response["json"]["aiInterpretation"]
    assert interpretation["summary"].startswith("当前数据量有限")
    assert interpretation["findings"][0]["evidenceRefs"]


def test_generic_mode_surfaces_provider_error_instead_of_summary_error(
    handler_request,
) -> None:
    client = ProviderErrorInterpretationClient()
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=_provider_run(client),
    ):
        response = handler_request({
            "filename": "monitoring.csv",
            "csvText": _generic_csv(),
            "providerMode": "openai_responses",
            "confirmLive": True,
            "confirmMaxCalls": 6,
        })
    assert response["status"] == 400
    assert "OpenAI 鉴权失败" in response["json"]["error"]
    assert "summary" not in response["json"]["error"]


def test_connection_probe_marks_provider_verified_without_user_data() -> None:
    runtime = DesktopProviderRuntime()
    provider_run = _provider_run(ConnectionProbeClient())
    with patch.object(runtime, "resolve", return_value=provider_run):
        result = probe_provider_connection(
            runtime, {"provider": "openai", "confirmProbe": True}
        )
    assert result["ok"] is True
    assert result["probeContainsUserData"] is False
    assert result["status"]["openai"]["connectionVerified"] is True


def test_connection_probe_hides_key_fragment_on_auth_failure() -> None:
    runtime = DesktopProviderRuntime()
    provider_run = _provider_run(ConnectionProbeClient(auth_error=True))
    with patch.object(runtime, "resolve", return_value=provider_run):
        with pytest.raises(DesktopProviderError) as exc_info:
            probe_provider_connection(
                runtime, {"provider": "openai", "confirmProbe": True}
            )
    message = str(exc_info.value)
    assert "OpenAI 鉴权失败" in message
    assert "sk-ws-" not in message
    assert "secret" not in message


def test_related_workbook_uses_provider_without_sending_raw_rows(
    handler_request,
) -> None:
    client = GroundedInterpretationClient()
    with patch(
        "cureflow_app.server.DesktopProviderRuntime.resolve",
        return_value=_provider_run(client),
    ):
        response = handler_request({
            "filename": "monitoring.xlsx",
            "xlsxBase64": _workbook_base64(),
            "providerMode": "openai_responses",
            "confirmLive": True,
            "confirmMaxCalls": 6,
        })
    assert response["status"] == 200
    payload = response["json"]
    assert payload["input"]["analysisMode"] == "workbook_profile"
    assert payload["input"]["cloudProviderUsed"] is True
    assert payload["cost"]["llmCalls"] == 1
    assert "2024-01-01" not in client.prompts[0]
    assert "2025-01-01" not in client.prompts[0]
    assert '"rawRowsSent": false' in client.prompts[0]


def test_generic_profile_rejects_more_than_one_hundred_columns() -> None:
    source = ",".join(f"c{i}" for i in range(101)) + "\n" + ",".join("1" for _ in range(101))
    with pytest.raises(ValueError, match="100 列"):
        prepare_tabular_csv(source)


def test_workbook_profile_preserves_training_experiment_relationship(tmp_path: Path) -> None:
    result = analyze_generic_workbook(_workbook_tables(), "monitoring.xlsx", tmp_path)
    schema = result["workbookSchema"]
    assert result["mode"] == "OFFLINE_DETERMINISTIC_WORKBOOK_PROFILE"
    assert schema["training_sheet"] == "训练集"
    assert schema["experiment_sheet"] == "实验集"
    assert schema["target_candidates"] == ["目标值"]
    assert schema["label_columns"] == ["阶段标签"]
    assert schema["feature_columns"] == ["降雨量", "孔隙水压力"]
    assert "目标值" in result["report"]
    assert "阶段标签" in result["report"]


def test_workbook_profile_persists_each_sheet_and_schema(tmp_path: Path) -> None:
    result = analyze_generic_workbook(_workbook_tables(), "monitoring.xlsx", tmp_path)
    run_dir = tmp_path / "artifacts" / result["runId"]
    assert (run_dir / "raw_sheet_1__v1.json").exists()
    assert (run_dir / "raw_sheet_2__v1.json").exists()
    assert (run_dir / "field_workbook_schema__v1.json").exists()


def test_api_routes_multi_sheet_xlsx_to_workbook_profile(handler_request) -> None:
    response = handler_request({"filename": "monitoring.xlsx", "xlsxBase64": _workbook_base64()})
    assert response["status"] == 200
    payload = response["json"]
    assert payload["input"]["analysisMode"] == "workbook_profile"
    assert payload["input"]["sheets"] == ["训练集", "实验集"]
    assert payload["workbookSchema"]["target_candidates"] == ["目标值"]
    assert payload["cost"]["llmCalls"] == 0
