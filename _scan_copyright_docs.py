from docx import Document
files = [
    r"C:\Users\Lenovo\Desktop\其他参与\ai\互联网+\软著资料\CureFlow软著登记页面填写稿V1.0.docx",
    r"C:\Users\Lenovo\Desktop\其他参与\ai\互联网+\软著资料\CureFlow软件说明书V1.0.docx",
    r"C:\Users\Lenovo\Desktop\其他参与\ai\互联网+\软著资料\CureFlow软件源代码V1.0.docx",
]
bad = ["842", "10 个", "10个", "13 页", "13页", "约 8", "8 千"]
ok  = ["8059", "20 个", "20个", "60 页", "60页", "3000 行", "截断"]
for f in files:
    doc = Document(f)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    print(f)
    print("  BAD hits:", [(w, text.count(w)) for w in bad if w in text])
    print("  OK  hits:", [(w, text.count(w)) for w in ok  if w in text])
    print()
