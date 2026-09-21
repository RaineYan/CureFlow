from docx import Document
import re
d = Document(r"C:\Users\Lenovo\Desktop\其他参与\ai\互联网+\软著资料\CureFlow软件说明书V1.0.docx")
print("--- paragraphs with desktop/electron/api ---")
for i, p in enumerate(d.paragraphs):
    if re.search(r"(desktop|electron|本地桌面|桌面端|沙箱|sandbox|api\s*/|vite|react|node\.?js)", p.text, re.IGNORECASE):
        print(f"p{i}: {p.text[:300]}")
print("--- tables ---")
for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if re.search(r"(desktop|electron|本地桌面|桌面端|沙箱|sandbox|api\s*/|vite|react)", cell.text, re.IGNORECASE):
                print(f"t{ti} r{ri} c{ci}: {cell.text[:200]}")
